import os
import io
import pandas as pd
import streamlit as st
import bigdataquery as bdq
from openpyxl.styles import Font, PatternFill, Alignment


# =========================================================
# 기본 설정
# =========================================================
TABLE_NAME = "dsehs_kh.dsehsmgr_t_sdq_acdt_rept_m"
DATE_COL = "ACDT_GNRT_DTM"
CREDENTIAL_DIR = "/config/.credential"

TARGET_DSRT_IDS = ["H1", "H2", "H3", "DSR"]

# 사고단계 코드 값별 색상 팔레트 (파스텔 배경 + 진한 글자)
STEP_CD_PALETTE = [
    "#E3F2FD",  # 파랑
    "#E8F5E9",  # 초록
    "#FFF3E0",  # 주황
    "#F3E5F5",  # 보라
    "#FFEBEE",  # 빨강
    "#E0F7FA",  # 청록
    "#FFF8E1",  # 노랑
    "#E8EAF6",  # 남색
    "#FCE4EC",  # 분홍
    "#F1F8E9",  # 연두
]
STEP_CD_TEXT_COLOR = "#1F1F1F"


# =========================================================
# 컬럼 정의
# source: DB 실제 컬럼명
# display: 화면/엑셀에 보여줄 한국어명
# =========================================================
COLUMN_META = [
    {"source": "ACDT_NO", "display": "사고번호"},
    {"source": "ACDT_STEP_CD", "display": "사고단계 코드"},
    {"source": "ACDT_STEP_DTLD_CD", "display": "사고단계 세부 코드"},
    {"source": "ACDT_NM", "display": "사고명"},
    {"source": "ACDT_GNRT_DTM", "display": "사고일시"},
    {"source": "ACDT_GRD_CD", "display": "사고 등급 코드"},
    {"source": "ACDT_CATG_CD", "display": "사고 범주 코드"},
    {"source": "ACDT_CLSF_CD", "display": "사고 분류 코드"},
    {"source": "ACDT_TYPE_CD", "display": "사고 유형 코드"},
    {"source": "SITE_ID", "display": "사업장ID"},
    {"source": "GBM_CD", "display": "사업부 코드"},
    {"source": "INCHG_DEPT_ID", "display": "책임부서ID(귀책 부서)"},
    {"source": "INCHG_COMP_ID", "display": "책임 업체 ID(귀책 협력사)"},
    {"source": "INCHG_COMP_CRNO", "display": "책임 업체 사업자등록번호"},
    {"source": "INCHG_BSPT_YN", "display": "귀책 협력사 유무"},
    {"source": "OR1_COMP_ID", "display": "1차 업체 ID"},
    {"source": "OR1_COMP_CRNO", "display": "1차 업체 사업자등록번호"},
    {"source": "HDQT_DSRT_ID", "display": "전자구역ID"},
    {"source": "HDQT_BUIL_ID", "display": "전사건물ID"},
    {"source": "HDQT_LCT_ID", "display": "전사위치ID"},
    {"source": "ACDT_BAY_VAL", "display": "사고 BAY 값"},
    {"source": "ACDT_PILAR_VAL", "display": "사고 기둥 값"},
    {"source": "ACDT_DTL_LCT_VAL", "display": "사고 상세 위치 값"},
    {"source": "ACDT_PLAC_CD", "display": "사고 장소 코드"},
    {"source": "ACDT_SDQ_TAT_CD", "display": "사고 TAT 코드"},
    {"source": "ACDT_DELY_DY_NUM", "display": "사고 지연 일수"},
    {"source": "FLFL_SDQ_TAT_CD", "display": "이행 TAT 코드"},
    {"source": "FLFL_DELY_DY_NUM", "display": "이행 지연 일수"},
    {"source": "FIREF_MV_YN", "display": "소방대 출동 여부"},
    {"source": "DCLR_RCIV_DTM", "display": "신고접수 일시"},
    {"source": "FIREF_NOT_DCLR_RSN", "display": "소방대 미신고 사유"},
    {"source": "EVCT_YN", "display": "대피여부"},
    {"source": "DISMNT_USR_ID", "display": "전파자 사용자 ID"},
    {"source": "DISMNT_USR_NM", "display": "전파자 사용자명"},
    {"source": "DISMNT_USR_NM_ENC", "display": "전파자 사용자명암호화"},
    {"source": "DISMNT_DTM", "display": "전파 일시"},
    {"source": "FIREF_RSPS_CNTN", "display": "소방대 대응 내용"},
    {"source": "GNRT_CNTN", "display": "발생 내용"},
    {"source": "FLFL_RSP_USR_ID", "display": "이행 담당 사용자 ID"},
    {"source": "FLFL_RSP_USR_NM", "display": "이행 담당 사용자명"},
    {"source": "FLFL_RSP_USR_NM_ENC", "display": "이행 담당 사용자명암호화"},
    {"source": "FNL_ACDT_CLSF_CD", "display": "최종 사고 분류 코드"},
    {"source": "FNL_ACDT_GRD_CD", "display": "최종 사고 등급 코드"},
    {"source": "CSTR_WRK_PRMSN_YN", "display": "공사 작업 허가 여부"},
    {"source": "CSTR_WRK_REQ_NO", "display": "공사작업의뢰번호"},
    {"source": "OCM_LCLSF_CD", "display": "기인물 대분류 코드"},
    {"source": "OCM_SCLSF_CD", "display": "기인물 소분류 코드"},
    {"source": "SDQ_RVLTN_CAUS_CD", "display": "발현원인코드"},
    {"source": "ACDT_SIT_CD", "display": "사고 상황 코드"},
    {"source": "ACDT_FCLT_INFO_CD", "display": "사고 설비 정보 코드"},
    {"source": "FCLT_MGT_DIV_CD", "display": "설비관리구분코드"},
    {"source": "FCLT_SITE_ID", "display": "설비 사업장 ID"},
    {"source": "MSTR_FCLT_ID", "display": "마스터 설비 ID"},
    {"source": "EQP_ID", "display": "EQPID"},
    {"source": "FCLT_MDL_NM", "display": "설비모델명"},
    {"source": "FCLT_INFO_DRCT_INP_CNTN", "display": "설비정보 직접 입력 내용"},
    {"source": "ACDT_FCLT_STT_CD", "display": "사고 설비 상태 코드"},
    {"source": "ACDT_REPT_PHOT_ATTH_DOC_ID", "display": "사고 보고 사진 첨부 문서 ID"},
    {"source": "ANYTM_RSKA_DIV_CD", "display": "수시위험성평가구분코드"},
    {"source": "SOP_INFO_VAL", "display": "SOP 정보 값"},
    {"source": "ANYTM_RSKA_ATTH_DOC_ID", "display": "수시위험성평가 첨부 문서 ID"},
    {"source": "EDU_FINS_DT", "display": "교육 납기 일자"},
    {"source": "EDU_DTL_CNTN", "display": "교육 상세 내용"},
]

SELECT_COLUMNS = [x["source"] for x in COLUMN_META]
DISPLAY_MAP = {x["source"]: x["display"] for x in COLUMN_META}


# =========================================================
# Streamlit 설정
# =========================================================
st.set_page_config(
    page_title="사고 리포트 조회",
    page_icon="🚨",
    layout="wide",
)

st.title("🚨 사고 리포트 조회")
st.caption("전자구역ID **H1 / H2 / H3 / DSR** 기준으로 전체 연도 사고 데이터를 조회합니다.")
st.divider()


# =========================================================
# 유틸 함수
# =========================================================
def token_exists(user_name: str) -> bool:
    token_path = f"{CREDENTIAL_DIR}/bigdataquery.token.{user_name}"
    return os.path.exists(token_path)


def format_dtm_14(value):
    """
    20260102153045 같은 14자리 문자열을
    2026-01-02 15:30:45 형태로 표시.
    원본 값이 형식에 안 맞으면 그대로 반환.
    """
    if pd.isna(value):
        return value

    value = str(value).strip()

    if len(value) == 14 and value.isdigit():
        return (
            f"{value[0:4]}-{value[4:6]}-{value[6:8]} "
            f"{value[8:10]}:{value[10:12]}:{value[12:14]}"
        )

    return value


def build_step_color_map(step_series: pd.Series) -> dict:
    """
    사고단계 코드 고유값마다 팔레트에서 색상을 하나씩 배정.
    값 순서를 정렬해서 재조회해도 같은 값 = 같은 색이 되게 함.
    """
    uniques = sorted(
        str(v).strip() for v in step_series.dropna().unique() if str(v).strip()
    )
    return {
        value: STEP_CD_PALETTE[i % len(STEP_CD_PALETTE)]
        for i, value in enumerate(uniques)
    }


def style_step_cd(display_df: pd.DataFrame, color_map: dict):
    """사고단계 코드 컬럼에 값별 배경색을 입힌 Styler 반환."""
    step_col = DISPLAY_MAP["ACDT_STEP_CD"]

    if step_col not in display_df.columns:
        return display_df

    def _color_cells(col):
        styles = []
        for v in col:
            key = "" if pd.isna(v) else str(v).strip()
            bg = color_map.get(key)
            if bg:
                styles.append(
                    f"background-color: {bg}; "
                    f"color: {STEP_CD_TEXT_COLOR}; "
                    f"font-weight: 600;"
                )
            else:
                styles.append("")
        return styles

    return display_df.style.apply(_color_cells, subset=[step_col])


def render_step_legend(color_map: dict):
    """사고단계 코드 색상 범례를 칩 형태로 표시."""
    if not color_map:
        return

    chips = "".join(
        f'<span style="display:inline-block; margin:2px 6px 2px 0; '
        f'padding:2px 10px; border-radius:12px; font-size:0.85rem; '
        f'font-weight:600; background-color:{bg}; color:{STEP_CD_TEXT_COLOR}; '
        f'border:1px solid rgba(0,0,0,0.1);">{value}</span>'
        for value, bg in color_map.items()
    )
    st.markdown(
        f'<div style="margin-bottom:8px;">'
        f'<span style="font-size:0.85rem; color:gray; margin-right:8px;">사고단계 코드 범례:</span>'
        f"{chips}</div>",
        unsafe_allow_html=True,
    )


@st.cache_data(ttl=60 * 10, show_spinner=False)
def load_acdt_data(user_name: str) -> pd.DataFrame:
    select_clause = ",\n    ".join(SELECT_COLUMNS)
    dsrt_list_sql = "', '".join(TARGET_DSRT_IDS)

    sql = f"""
SELECT
    {select_clause}
FROM {TABLE_NAME}
WHERE UPPER(TRIM(HDQT_DSRT_ID)) IN ('{dsrt_list_sql}')
ORDER BY {DATE_COL}
"""

    df = bdq.getData(sql, user_name=user_name)

    # bigdataquery 결과 컬럼이 소문자로 올 수 있어서 대문자로 통일
    df.columns = [str(col).upper() for col in df.columns]

    # 정의한 컬럼 순서대로 정렬
    existing_cols = [col for col in SELECT_COLUMNS if col in df.columns]
    df = df[existing_cols]

    # 전자구역ID 정리
    if "HDQT_DSRT_ID" in df.columns:
        df["HDQT_DSRT_ID"] = (
            df["HDQT_DSRT_ID"]
            .astype("string")
            .str.strip()
            .str.upper()
        )

    # 일시 컬럼 보기 좋게 변환
    for dt_col in ["ACDT_GNRT_DTM", "DCLR_RCIV_DTM", "DISMNT_DTM"]:
        if dt_col in df.columns:
            df[dt_col] = df[dt_col].apply(format_dtm_14)

    return df


def make_excel_bytes(df: pd.DataFrame) -> bytes:
    output = io.BytesIO()

    # 엑셀에는 한국어 컬럼명으로 저장
    display_df = df.rename(columns=DISPLAY_MAP)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        display_df.to_excel(writer, index=False, sheet_name="사고리포트")

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")

        ws = writer.sheets["사고리포트"]

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for column_cells in ws.columns:
            column_letter = column_cells[0].column_letter
            max_length = 0

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 45)

    output.seek(0)
    return output.getvalue()


def make_csv_bytes(df: pd.DataFrame) -> bytes:
    display_df = df.rename(columns=DISPLAY_MAP)
    return display_df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


# =========================================================
# 사이드바
# =========================================================
st.sidebar.header("⚙️ 조회 설정")

user_name = st.sidebar.text_input("SSO ID", value="dongrami.kim")

st.sidebar.divider()

with st.sidebar.expander("📋 조회 테이블 / 조건", expanded=False):
    st.write("조회 테이블")
    st.code(TABLE_NAME)
    st.write("조회 조건")
    st.code(
        "HDQT_DSRT_ID IN ('H1', 'H2', 'H3', 'DSR')\n연도 조건 없음",
        language="sql",
    )

show_sql = st.sidebar.checkbox("SQL 보기", value=False)

if show_sql:
    st.sidebar.code(
        f"""
SELECT
    {", ".join(SELECT_COLUMNS)}
FROM {TABLE_NAME}
WHERE UPPER(TRIM(HDQT_DSRT_ID)) IN ('H1', 'H2', 'H3', 'DSR')
ORDER BY {DATE_COL}
""",
        language="sql",
    )


# =========================================================
# 데이터 조회
# =========================================================
st.subheader("데이터 조회")

if "acdt_df" not in st.session_state:
    st.session_state["acdt_df"] = None

if st.button("🔍 전체 연도 데이터 조회", type="primary"):
    if not user_name:
        st.error("SSO ID를 입력하세요.")
        st.stop()

    if not token_exists(user_name):
        st.error("bigdataquery 토큰이 없습니다. 토큰 발급 후 다시 조회하세요.")
        st.stop()

    with st.spinner("사고 데이터를 조회 중입니다..."):
        try:
            df = load_acdt_data(user_name)
            st.session_state["acdt_df"] = df
        except Exception as e:
            st.error("데이터 조회 실패")
            st.exception(e)
            st.stop()


# =========================================================
# 조회 결과 표시 및 다운로드
# =========================================================
df = st.session_state.get("acdt_df")

if df is not None:
    st.success("조회 완료")

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("조회 건수", f"{len(df):,}건")
    col2.metric("컬럼 수", f"{len(df.columns):,}개")
    col3.metric(
        "사고단계 코드 종류",
        f"{df['ACDT_STEP_CD'].nunique():,}개" if "ACDT_STEP_CD" in df.columns else "-",
    )
    col4.metric("전자구역ID", "H1 / H2 / H3 / DSR")

    st.divider()

    # -----------------------------------------------------
    # 결과 필터 (화면 표시용 — 원본 데이터는 그대로 유지)
    # -----------------------------------------------------
    st.subheader("조회 결과")

    fcol1, fcol2, fcol3 = st.columns([1, 1, 2])

    with fcol1:
        dsrt_options = (
            sorted(df["HDQT_DSRT_ID"].dropna().unique())
            if "HDQT_DSRT_ID" in df.columns
            else []
        )
        dsrt_filter = st.multiselect(
            "전자구역ID 필터",
            options=dsrt_options,
            placeholder="전체",
        )

    with fcol2:
        step_options = (
            sorted(str(v).strip() for v in df["ACDT_STEP_CD"].dropna().unique())
            if "ACDT_STEP_CD" in df.columns
            else []
        )
        step_filter = st.multiselect(
            "사고단계 코드 필터",
            options=step_options,
            placeholder="전체",
        )

    with fcol3:
        keyword = st.text_input(
            "사고명 검색",
            placeholder="사고명에 포함된 키워드 입력",
        )

    view_df = df

    if dsrt_filter and "HDQT_DSRT_ID" in view_df.columns:
        view_df = view_df[view_df["HDQT_DSRT_ID"].isin(dsrt_filter)]

    if step_filter and "ACDT_STEP_CD" in view_df.columns:
        view_df = view_df[
            view_df["ACDT_STEP_CD"].astype(str).str.strip().isin(step_filter)
        ]

    if keyword and "ACDT_NM" in view_df.columns:
        view_df = view_df[
            view_df["ACDT_NM"]
            .astype(str)
            .str.contains(keyword, case=False, na=False)
        ]

    st.caption(f"표시 중: **{len(view_df):,}건** / 전체 {len(df):,}건")

    # -----------------------------------------------------
    # 사고단계 코드 색상 범례 + 테이블
    # -----------------------------------------------------
    step_color_map = (
        build_step_color_map(df["ACDT_STEP_CD"])
        if "ACDT_STEP_CD" in df.columns
        else {}
    )

    render_step_legend(step_color_map)

    display_df = view_df.rename(columns=DISPLAY_MAP)
    styled_df = style_step_cd(display_df, step_color_map)

    st.dataframe(
        styled_df,
        use_container_width=True,
        height=600,
    )

    st.divider()

    # -----------------------------------------------------
    # 다운로드 (현재 필터가 적용된 데이터 기준)
    # -----------------------------------------------------
    st.subheader("다운로드")
    st.caption("현재 화면에 표시된(필터 적용된) 데이터를 다운로드합니다.")

    csv_data = make_csv_bytes(view_df)
    excel_data = make_excel_bytes(view_df)

    dcol1, dcol2 = st.columns(2)

    with dcol1:
        st.download_button(
            label="⬇️ CSV 다운로드",
            data=csv_data,
            file_name="acdt_rept_all_years_h1_h2_h3_dsr.csv",
            mime="text/csv",
            use_container_width=True,
        )

    with dcol2:
        st.download_button(
            label="⬇️ 엑셀 다운로드",
            data=excel_data,
            file_name="acdt_rept_all_years_h1_h2_h3_dsr.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
else:
    st.info("`전체 연도 데이터 조회` 버튼을 누르면 데이터를 가져옵니다.")
